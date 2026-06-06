import streamlit as st
from io import BytesIO
from pypdf import PdfReader

# Securely import modern Google GenAI SDK
try:
    from google import genai
except ImportError:
    st.error("The 'google-genai' package is missing. Please run 'pip install google-genai' in your VS Code terminal.")

# Import Document generation tools
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# ---------------------------------------------------------------------------
# CORE APPLICATION LOGIC
# ---------------------------------------------------------------------------

def extract_text_from_pdf(uploaded_file):
    """Safely extracts text content from an uploaded resume PDF."""
    try:
        pdf_reader = PdfReader(uploaded_file)
        text = ""
        for page in pdf_reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        return text
    except Exception as e:
        st.error(f"Error parsing PDF file: {e}")
        return ""

def analyze_resume_with_ai(api_key, resume_text, job_description):
    """Connects to Gemini 2.5 Flash API in real-time for ATS evaluation."""
    try:
        client = genai.Client(api_key=api_key)
        prompt = f"""
        You are an expert ATS (Applicant Tracking System) scanner and senior corporate recruiter. 
        Analyze the following Resume against the Job Description.
        
        Provide your evaluation in these exact sections:
        1. OVERALL ATS MATCH SCORE (Out of 100)
        2. CRITICAL GAPS IDENTIFIED (Missing skills, keywords, or background)
        3. ACTIONABLE TAILORING RECOMMENDATIONS (Specific adjustments to fix gaps)
        4. SUGGESTED OPTIMIZED PROFESSIONAL SUMMARY (Ready to copy-paste)
        
        Target Job Description:
        {job_description}
        
        Applicant Resume:
        {resume_text}
        """
        response = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=prompt
        )
        return response.text
    except Exception as e:
        return f"AI Connection Error: {e}. Please check if your Gemini API key is correct."

# ---------------------------------------------------------------------------
# DOWNLOADABLE FILE GENERATORS
# ---------------------------------------------------------------------------

def generate_docx(report_text):
    """Compiles the AI feedback into a downloadable Word file."""
    doc = Document()
    doc.add_heading('ATS Resume Optimization Report', level=1)
    for line in report_text.split('\n'):
        if line.strip():
            doc.add_paragraph(line)
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

def generate_pdf(report_text):
    """Compiles the AI feedback into a cleanly formatted PDF file."""
    bio = BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [Paragraph("<b>ATS Optimization Report</b>", styles['Title']), Spacer(1, 15)]
    
    for line in report_text.split('\n'):
        if line.strip():
            # Clean basic symbols for PDF library compatibility
            clean_line = line.replace('*', '').replace('#', '').strip()
            story.append(Paragraph(clean_line, styles['Normal']))
            story.append(Spacer(1, 6))
            
    doc.build(story)
    bio.seek(0)
    return bio

# ---------------------------------------------------------------------------
# USER INTERFACE DESIGN
# ---------------------------------------------------------------------------

st.set_page_config(page_title="AI Resume ATS Optimizer", page_icon="🎯", layout="wide")

st.title("🎯 AI Resume ATS Ranking & Gap Optimizer")
st.write("Upload your current CV and paste your target job profile to get instantaneous scoring, benchmark tracking, and document exports.")

# Sidebar setup for security keys
api_key_input = st.sidebar.text_input("Enter Gemini API Key:", type="password", help="Get your key from Google AI Studio")

# Grid layout split for workspace efficiency
col1, col2 = st.columns(2)

with col1:
    st.subheader("📋 1. Target Job Parameters")
    job_description = st.text_area(
        "Paste the Target Job Description here:", 
        height=280, 
        placeholder="Paste requirements, core duties, or stack qualifications..."
    )

with col2:
    st.subheader("📄 2. Your Current Document")
    uploaded_file = st.file_uploader(
        "Upload your current Resume/CV (PDF Format):", 
        type=["pdf"]
    )

st.markdown("---")

# Execution trigger
if st.button("Analyze Alignment Gaps", type="primary", use_container_width=True):
    if not api_key_input:
        st.warning("Please input your Gemini API Key in the left sidebar layout pane.")
    elif not job_description:
        st.warning("Please supply a targeted Job Description to compare against.")
    elif not uploaded_file:
        st.warning("Please upload a PDF version of your resume.")
    else:
        with st.spinner("Analyzing resume against role requirements in real time..."):
            resume_text = extract_text_from_pdf(uploaded_file)
            if resume_text:
                ai_report = analyze_resume_with_ai(api_key_input, resume_text, job_description)
                st.session_state['active_report'] = ai_report

# Handle persistence and downloads
if 'active_report' in st.session_state:
    st.subheader("📊 Live ATS Matrix & Gap Analytics")
    st.markdown(st.session_state['active_report'])
    
    st.markdown("---")
    st.subheader("📥 Export Structural Alterations")
    
    dl_word, dl_pdf = st.columns(2)
    with dl_word:
        st.download_button(
            label="Download Recommendations as Word (.docx)",
            data=generate_docx(st.session_state['active_report']),
            file_name="ATS_Resume_Optimization_Guide.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
    with dl_pdf:
        st.download_button(
            label="Download Recommendations as PDF (.pdf)",
            data=generate_pdf(st.session_state['active_report']),
            file_name="ATS_Resume_Optimization_Guide.pdf",
            mime="application/pdf",
            use_container_width=True
        )